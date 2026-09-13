from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .layout_ir import PageLayout
from .oxml_shapes import (
    append_drawing_to_paragraph,
    ensure_cjk_doc_fonts,
    make_framed_text_paragraph,
    make_picture_drawing,
    make_rect_drawing,
    register_png,
)
from .renderer import PageCapture

PAGE_W_MM = 210.0
PAGE_H_MM = 297.0


def _clear_body(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _ensure_drawing_namespaces(doc: Document) -> None:
    """Declare DrawingML / Word 2010 shape namespaces on document root."""
    root = doc.element
    ns = {
        "xmlns:wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "xmlns:a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "xmlns:pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
        "xmlns:wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
        "xmlns:mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    }
    for key, val in ns.items():
        if key not in root.attrib.values() and root.get(key) is None:
            # set via Clark or string attr used by lxml
            try:
                root.set(key, val)
            except Exception:  # noqa: BLE001
                pass
    # mc:Ignorable for wps
    ign = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"
    existing = root.get(ign) or ""
    if "wps" not in existing.split():
        root.set(ign, (existing + " wps").strip())


def _configure_section(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(PAGE_W_MM)
    section.page_height = Mm(PAGE_H_MM)
    section.left_margin = Mm(0)
    section.right_margin = Mm(0)
    section.top_margin = Mm(0)
    section.bottom_margin = Mm(0)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)


def _blank_paragraph(doc: Document):
    paragraph = doc.add_paragraph()
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    return paragraph


def _estimate_textbox_height_mm(
    text: str,
    w_mm: float,
    font_pt: float,
    line_height: float,
) -> float:
    """Estimate Word/WPS wrap height for CJK-heavy text."""
    char_w_mm = max(font_pt * 25.4 / 72 * 1.08, 0.8)
    chars_per_line = max(1, int(w_mm / char_w_mm))
    lines = 0
    for part in text.split("\n") or [text]:
        part = part.strip() or " "
        lines += max(1, (len(part) + chars_per_line - 1) // chars_per_line)
    line_h_mm = font_pt * 25.4 / 72 * max(line_height, 1.15)
    return lines * line_h_mm + 1.0


def _paint_layout(doc: Document, layout: PageLayout, shape_id_start: int) -> int:
    sid = shape_id_start
    paragraph = _blank_paragraph(doc)

    for rect in layout.rects:
        drawing = make_rect_drawing(
            x_mm=rect.x_mm,
            y_mm=rect.y_mm,
            w_mm=rect.w_mm,
            h_mm=max(rect.h_mm, 0.15),
            fill_rgb=rect.fill_rgb,
            fill_alpha=rect.fill_alpha,
            shape_id=sid,
        )
        append_drawing_to_paragraph(paragraph, drawing)
        sid += 1

    for image in layout.images:
        if not image.png:
            continue
        r_id = register_png(doc, image.png)
        drawing = make_picture_drawing(
            r_id=r_id,
            x_mm=image.x_mm,
            y_mm=image.y_mm,
            w_mm=image.w_mm,
            h_mm=image.h_mm,
            shape_id=sid,
        )
        append_drawing_to_paragraph(paragraph, drawing)
        sid += 1

    # 文字用 framePr 普通段落，避免 DrawingML 文本框在 WPS/部分预览中文方块
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    for text in layout.texts:
        # 宽段落少加水平 padding，避免右侧栏文字越界换行错位
        pad_x = max(0.8, text.w_mm * (0.04 if text.w_mm >= 60 else 0.10))
        pad_y = max(0.3, text.h_mm * 0.05)
        font_pt = max(5.0, text.font_size_pt * 0.92)
        x_mm = max(0.0, text.x_mm)
        y_mm = max(0.0, text.y_mm)
        w_mm = min(PAGE_W_MM - x_mm, max(text.w_mm + pad_x, 4.0))
        base_h = text.h_mm + pad_y
        est_h = _estimate_textbox_height_mm(text.text, w_mm, font_pt, text.line_height)
        if est_h > base_h:
            h_mm = min(est_h, text.h_mm * 1.15 + pad_y)
        else:
            h_mm = base_h
        h_mm = min(PAGE_H_MM - y_mm, max(h_mm, 2.5))
        lh = min(max(text.line_height, 1.08), 1.35)
        framed = make_framed_text_paragraph(
            text=text.text,
            x_mm=x_mm,
            y_mm=y_mm,
            w_mm=w_mm,
            h_mm=h_mm,
            font_size_pt=font_pt,
            bold=text.bold,
            color_rgb=text.color_rgb,
            align=text.align,
            line_height=lh,
        )
        if sect is not None:
            sect.addprevious(framed)
        else:
            body.append(framed)

    return sid


def build_editable_docx(pages: list[PageCapture]) -> bytes:
    """按 LayoutIR 绝对定位铺版，生成可编辑且布局贴近 HTML 的 Word。"""
    doc = Document()
    _clear_body(doc)
    _ensure_drawing_namespaces(doc)
    _configure_section(doc)
    ensure_cjk_doc_fonts(doc)

    shape_id = 1
    for i, page in enumerate(pages):
        shape_id = _paint_layout(doc, page.layout, shape_id)
        if i < len(pages) - 1:
            # page break on a dedicated empty paragraph
            br = _blank_paragraph(doc)
            run = br.add_run()
            run.add_break(WD_BREAK.PAGE)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
