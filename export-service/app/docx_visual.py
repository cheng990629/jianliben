from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image

from .renderer import PageCapture

# 略小于 A4，避免图片 + 段落行距溢出产生空白页
PAGE_W_MM = 210.0
PAGE_H_MM = 297.0
FIT_W_MM = 209.0
FIT_H_MM = 296.0


def _clear_body(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _fit_mm(png: bytes) -> tuple[float, float]:
    """按截图像素比缩进 FIT_W×FIT_H 框内，保证单页放得下。"""
    with Image.open(io.BytesIO(png)) as im:
        w, h = im.size
    if w <= 0 or h <= 0:
        return FIT_W_MM, FIT_H_MM
    aspect = w / h
    box_aspect = FIT_W_MM / FIT_H_MM
    if aspect >= box_aspect:
        width = FIT_W_MM
        height = FIT_W_MM / aspect
    else:
        height = FIT_H_MM
        width = FIT_H_MM * aspect
    return width, height


def build_visual_docx(pages: list[PageCapture]) -> bytes:
    """每页截图嵌入 Word，版式与 HTML A4 视觉一致（不可改字）。"""
    doc = Document()
    _clear_body(doc)

    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(PAGE_W_MM)
    section.page_height = Mm(PAGE_H_MM)
    section.left_margin = Mm(0)
    section.right_margin = Mm(0)
    section.top_margin = Mm(0)
    section.bottom_margin = Mm(0)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)

    for i, page in enumerate(pages):
        width_mm, height_mm = _fit_mm(page.png)
        paragraph = doc.add_paragraph()
        pf = paragraph.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.line_spacing = 1.0

        run = paragraph.add_run()
        run.add_picture(
            io.BytesIO(page.png),
            width=Mm(width_mm),
            height=Mm(height_mm),
        )
        # 分页符放在同一段末尾，避免「整页图 + 独立分页段」再挤出空白页
        if i < len(pages) - 1:
            run.add_break(WD_BREAK.PAGE)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
